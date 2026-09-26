"""Proposta comercial: leitura das regras da proposta no edital, formação de preço (custo + BDI),
checagem de exequibilidade, minuta pela IA e exportação em Word."""
import io
import re
from datetime import date

from flask import current_app

from extensions import db
from services import kb, llm, tabelas
from services.fluxos import cortar
from services.prompts import BASE, _j

# ---------------------------------------------------------------- regimes tributários (sugestões editáveis)
# Alíquota aproximada sobre o preço de venda, para o cálculo do BDI. O usuário sempre pode ajustar.
REGIMES = {
    "simples": {"nome": "Simples Nacional", "tributos_pct": 6.0,
                "nota": "Use a alíquota efetiva do seu anexo e faixa do Simples, excluída a parcela de IRPJ/CSLL (consulte o contador)."},
    "presumido": {"nome": "Lucro Presumido", "tributos_pct": 8.65,
                  "nota": "PIS 0,65% + COFINS 3% + ISS 5% (ajuste o ISS do município). IRPJ e CSLL não entram no BDI "
                          "(orientação do TCU): considere-os no lucro."},
    "real": {"nome": "Lucro Real", "tributos_pct": 14.25,
             "nota": "PIS 1,65% + COFINS 7,6% (não cumulativos) + ISS 5%. IRPJ e CSLL incidem sobre o lucro e não entram no BDI."},
}


def bdi_tcu(ac, sg, r, df, l, i):
    """Fórmula do Acórdão TCU 2.622/2013: BDI = [(1+AC+S+R+G)(1+DF)(1+L)/(1−I)] − 1 (valores em %)."""
    ac, sg, r, df, l, i = (float(x or 0) / 100 for x in (ac, sg, r, df, l, i))
    if i >= 1:
        return None
    return round((((1 + ac + sg + r) * (1 + df) * (1 + l)) / (1 - i) - 1) * 100, 2)


def _num(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return tabelas.para_preco(v)


def recalcular(proposta):
    """Recalcula preço unitário/total de cada item e os alertas objetivos (regras, sem IA)."""
    obra = bool((proposta.condicoes or {}).get("obra_ou_servico_engenharia"))
    itens, alertas = [], []
    total = total_estimado = 0.0
    for n, it in enumerate(proposta.itens or [], start=1):
        it = dict(it)
        qtd = _num(it.get("quantidade")) or 0
        custo = _num(it.get("custo_unitario"))
        bdi = _num(it.get("bdi"))
        bdi = proposta.bdi_pct if bdi is None else bdi
        if it.get("preco_manual") and _num(it.get("preco_unitario")) is not None:
            preco = round(_num(it["preco_unitario"]), 2)
        elif custo is not None:
            preco = round(custo * (1 + (bdi or 0) / 100), 2)
        else:
            preco = None
        it.update({"quantidade": qtd, "custo_unitario": custo, "bdi": _num(it.get("bdi")), "preco_unitario": preco,
                   "preco_total": round(preco * qtd, 2) if preco is not None else None})
        est = _num(it.get("valor_unitario_estimado"))
        rotulo = f"Item {it.get('numero') or n}"
        situacao = []
        if preco is not None:
            total += preco * qtd
            if est:
                total_estimado += est * qtd
                razao = preco / est
                it["pct_estimado"] = round(razao * 100, 1)
                if razao > 1:
                    situacao.append("acima_estimado")
                    alertas.append({"nivel": "alto", "texto": f"{rotulo}: preço {razao * 100:.0f}% do estimado — acima do "
                                    "orçamento da Administração, sujeito a desclassificação.", "fundamento": "Lei 14.133, art. 59, III"})
                elif obra and razao < 0.75:
                    situacao.append("inexequivel")
                    alertas.append({"nivel": "alto", "texto": f"{rotulo}: abaixo de 75% do orçado — presunção de "
                                    "inexequibilidade em obras e serviços de engenharia.", "fundamento": "Lei 14.133, art. 59, §4º"})
                elif obra and razao < 0.85:
                    situacao.append("garantia_adicional")
                    alertas.append({"nivel": "medio", "texto": f"{rotulo}: abaixo de 85% do orçado — será exigida "
                                    "garantia adicional se vencer.", "fundamento": "Lei 14.133, art. 59, §5º"})
                elif not obra and razao < 0.5:
                    situacao.append("indicio_inexequivel")
                    alertas.append({"nivel": "medio", "texto": f"{rotulo}: abaixo de 50% do orçado — indício de "
                                    "inexequibilidade; prepare a demonstração de custos.",
                                    "fundamento": "IN SEGES/ME 73/2022, art. 34"})
            if custo is not None and preco < custo:
                situacao.append("prejuizo")
                alertas.append({"nivel": "alto", "texto": f"{rotulo}: preço abaixo do custo informado.", "fundamento": ""})
        mediana = ((it.get("referencias") or {}).get("mercado") or {}).get("mediana")
        if preco is not None and mediana:
            it["pct_mercado"] = round(preco / mediana * 100, 1)
        it["situacao"] = situacao
        itens.append(it)
    proposta.itens = itens
    proposta.alertas = [a for a in (proposta.alertas or []) if a.get("origem") == "ia"] + \
        [{**a, "origem": "regra"} for a in alertas]
    return {"total": round(total, 2), "total_estimado": round(total_estimado, 2) if total_estimado else None}


def totais(proposta):
    total = sum((it.get("preco_total") or 0) for it in proposta.itens or [])
    custo = sum((_num(it.get("custo_unitario")) or 0) * (_num(it.get("quantidade")) or 0) for it in proposta.itens or [])
    est = sum((_num(it.get("valor_unitario_estimado")) or 0) * (_num(it.get("quantidade")) or 0)
              for it in proposta.itens or [])
    return {"total": round(total, 2), "custo_total": round(custo, 2), "total_estimado": round(est, 2) or None,
            "margem_bruta": round(total - custo, 2) if custo else None}


# ---------------------------------------------------------------- leitura do edital (IA barata)
def prompt_condicoes(texto):
    sistema = BASE + " Sua tarefa agora é apenas EXTRAIR do edital o que se refere à PROPOSTA COMERCIAL, sem opinar."
    usuario = f"""Extraia do edital abaixo, em JSON:
{{
 "objeto": "", "orgao": "", "numero": "", "criterio_julgamento": "", "modo_disputa": "",
 "obra_ou_servico_engenharia": false, "valor_estimado_global": null, "orcamento_sigiloso": false,
 "itens": [{{"numero": "", "lote": "", "descricao": "", "unidade": "", "quantidade": null,
            "valor_unitario_estimado": null, "codigo_catmat_catser": "", "pagina": ""}}],
 "validade_minima_dias": null, "prazo_entrega_execucao": "", "local_entrega": "", "condicoes_pagamento": "",
 "reajuste": "", "exige_planilha_custos": false, "exige_composicao_bdi": false, "bdi_referencia": "",
 "exige_marca_modelo": false, "amostra_ou_prova_conceito": "", "garantia_proposta": "",
 "casas_decimais": "", "forma_apresentacao": "", "documentos_com_proposta": [""],
 "declaracoes_exigidas": [""], "observacoes_importantes": [""]
}}
Regras: liste TODOS os itens/lotes da planilha do termo de referência (até 200), com quantidades e valores
estimados unitários quando o edital os divulgar (use null se o orçamento for sigiloso ou não constar).
Números sem símbolo de moeda, com ponto decimal. Indique a página no formato do texto ([pág. N]).
"declaracoes_exigidas": declarações que o edital exige junto com a proposta (ex.: que os preços incluem
todos os custos, cumprimento de reserva de cargos, enquadramento ME/EPP). Não invente o que não constar.

EDITAL:
{texto}"""
    demo = {
        "objeto": "Prestação de serviços contínuos de limpeza e conservação predial", "orgao": "Prefeitura Municipal de Exemplo",
        "numero": "Pregão Eletrônico 45/2026", "criterio_julgamento": "Menor preço global", "modo_disputa": "Aberto",
        "obra_ou_servico_engenharia": False, "valor_estimado_global": 1850000.0, "orcamento_sigiloso": False,
        "itens": [
            {"numero": "1", "lote": "", "descricao": "Servente de limpeza — 44h semanais (posto)", "unidade": "posto/mês",
             "quantidade": 24, "valor_unitario_estimado": 5600.0, "codigo_catmat_catser": "", "pagina": "38"},
            {"numero": "2", "lote": "", "descricao": "Encarregado de limpeza — 44h semanais", "unidade": "posto/mês",
             "quantidade": 2, "valor_unitario_estimado": 7200.0, "codigo_catmat_catser": "", "pagina": "38"},
            {"numero": "3", "lote": "", "descricao": "Material de limpeza e higiene (fornecimento mensal)", "unidade": "mês",
             "quantidade": 12, "valor_unitario_estimado": 3900.0, "codigo_catmat_catser": "", "pagina": "39"},
        ],
        "validade_minima_dias": 60, "prazo_entrega_execucao": "12 meses, prorrogável", "local_entrega": "Prédios da Prefeitura",
        "condicoes_pagamento": "Até 30 dias após o atesto da nota fiscal", "reajuste": "Repactuação anual (CCT) e IPCA",
        "exige_planilha_custos": True, "exige_composicao_bdi": False, "bdi_referencia": "",
        "exige_marca_modelo": False, "amostra_ou_prova_conceito": "", "garantia_proposta": "Não exigida",
        "casas_decimais": "2 casas decimais", "forma_apresentacao": "Proposta readequada ao lance final em até 2 horas",
        "documentos_com_proposta": ["Planilha de custos e formação de preços (modelo IN 05/2017)", "Indicação da CCT utilizada"],
        "declaracoes_exigidas": ["Que nos preços estão incluídos todos os custos, tributos e encargos",
                                 "Enquadramento como ME/EPP, se for o caso"],
        "observacoes_importantes": ["Salários não podem ser inferiores ao piso da CCT indicada no edital (pág. 41)"],
    }
    return sistema, usuario, demo


def ler_condicoes(proposta, edital):
    s, u, demo = prompt_condicoes(cortar(edital.texto))
    r = llm.chamar("barata", s, u, max_tokens=8000, demo=demo)
    cond = llm.extrair_json(r.texto)
    itens = []
    for it in (cond.get("itens") or [])[:200]:
        if not str(it.get("descricao") or "").strip():
            continue
        itens.append({"numero": str(it.get("numero") or ""), "lote": str(it.get("lote") or ""),
                      "descricao": str(it["descricao"]).strip(), "unidade": str(it.get("unidade") or ""),
                      "quantidade": _num(it.get("quantidade")) or 1,
                      "valor_unitario_estimado": _num(it.get("valor_unitario_estimado")),
                      "catmat": str(it.get("codigo_catmat_catser") or "").strip(), "pagina": str(it.get("pagina") or ""),
                      "custo_unitario": None, "bdi": None})
    cond.pop("itens", None)
    proposta.condicoes = cond
    proposta.itens = itens
    if cond.get("validade_minima_dias"):
        try:
            proposta.validade_dias = max(int(cond["validade_minima_dias"]), proposta.validade_dias or 0)
        except (TypeError, ValueError):
            pass
    proposta.demonstracao = r.demonstracao
    return [r]


# ---------------------------------------------------------------- minuta (IA de redação)
def _cnpj(c):
    c = re.sub(r"\D", "", c or "")
    return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}" if len(c) == 14 else c


def prompt_minuta(proposta, empresa, edital, tot, refs, contexto):
    cond = proposta.condicoes or {}
    itens_resumo = [{"numero": it.get("numero"), "descricao": it.get("descricao"), "unidade": it.get("unidade"),
                     "quantidade": it.get("quantidade"), "preco_unitario": it.get("preco_unitario"),
                     "valor_unitario_estimado": it.get("valor_unitario_estimado"),
                     "mediana_mercado": ((it.get("referencias") or {}).get("mercado") or {}).get("mediana"),
                     "custo_unitario": it.get("custo_unitario")} for it in (proposta.itens or [])[:60]]
    sistema = BASE + (" Você redige PROPOSTAS COMERCIAIS de empresas licitantes e revisa a formação de preço. "
                      "A tabela de itens e o valor por extenso são inseridos pelo sistema: no texto, escreva apenas o "
                      "marcador {{TABELA_ITENS}} na posição da tabela e {{VALOR_TOTAL}} onde couber o valor global. "
                      "Inclua somente declarações e documentos que o edital exige (lista fornecida) e as condições "
                      "informadas; não invente cláusulas, prazos nem fundamentos. Use [preencher] para dados ausentes.")
    usuario = f"""EMPRESA LICITANTE: {_j({"razao_social": empresa.razao_social, "cnpj": _cnpj(empresa.cnpj), "porte": empresa.porte})}
EDITAL: {_j({"orgao": cond.get("orgao") or (edital.orgao if edital else None), "numero": cond.get("numero") or (edital.numero if edital else None),
              "objeto": cond.get("objeto") or (edital.objeto if edital else None), "modalidade": edital.modalidade if edital else None})}
O QUE O EDITAL EXIGE DA PROPOSTA: {_j(cond)}
PARÂMETROS DA PROPOSTA: {_j({"validade_dias": proposta.validade_dias, "regime_tributario": REGIMES.get(proposta.regime, {}).get("nome"),
                              "bdi_pct": proposta.bdi_pct, "tributos_pct": proposta.tributos_pct, **tot})}
ITENS PRECIFICADOS: {_j(itens_resumo)}
REFERÊNCIAS DAS TABELAS OFICIAIS (quando encontradas):
{refs or "nenhuma tabela de referência compatível"}
BASE LEGAL DE APOIO:
{contexto}

Devolva JSON:
{{"texto": "minuta completa da proposta comercial em markdown simples (## para títulos), dirigida ao órgão, com: identificação da licitante e do edital, objeto, {{{{TABELA_ITENS}}}}, valor global {{{{VALOR_TOTAL}}}}, validade da proposta, prazo e local de execução/entrega, condições de pagamento, declarações exigidas, local, data e assinatura do representante legal",
 "alertas": [{{"nivel": "alto|medio|baixo", "texto": "ponto de atenção sobre preço, exequibilidade ou exigência da proposta", "fundamento": "artigo/norma ou vazio"}}]}}
Nos alertas, compare os preços com o estimado, com a mediana de mercado e com as tabelas oficiais; aponte
documentos obrigatórios que faltam (ex.: planilha de custos, composição do BDI) e riscos de desclassificação."""
    demo = {
        "texto": ("## Proposta comercial\n\nÀ **{orgao}**\nRef.: {numero}\n\n**{razao}**, inscrita no CNPJ sob o nº {cnpj}, "
                  "apresenta sua proposta para {objeto}, nos termos do edital.\n\n## Itens e preços\n\n{{{{TABELA_ITENS}}}}\n\n"
                  "**Valor global:** {{{{VALOR_TOTAL}}}}\n\n## Condições\n\n- Validade da proposta: {validade} dias.\n"
                  "- Prazo de execução: [preencher].\n- Pagamento: conforme o edital.\n\n## Declarações\n\n"
                  "Declaramos que nos preços propostos estão incluídos todos os custos, tributos, encargos sociais e "
                  "trabalhistas e demais despesas necessárias à execução do objeto.\n\n[Local], [data].\n\n"
                  "______________________________\nRepresentante legal").format(
            orgao=cond.get("orgao") or "[órgão]", numero=cond.get("numero") or "[número do edital]",
            razao=empresa.razao_social, cnpj=_cnpj(empresa.cnpj), objeto=cond.get("objeto") or "[objeto]",
            validade=proposta.validade_dias),
        "alertas": [{"nivel": "medio", "texto": "Proposta de demonstração: configure as chaves de IA para a revisão real dos preços.",
                     "fundamento": ""}],
    }
    return sistema, usuario, demo


def gerar_minuta(proposta, empresa, edital):
    tot = totais(proposta)
    refs = tabelas.contexto_para_ia(proposta.itens or [], uf=(empresa.ufs or "")[:2] or None)
    contexto = kb.buscar("proposta preço exequibilidade inexequível planilha custos BDI tabela referência "
                         "SINAPI convenção coletiva encargos tributos validade declaração " +
                         ((proposta.condicoes or {}).get("objeto") or ""), k=6)
    s, u, demo = prompt_minuta(proposta, empresa, edital, tot, refs, contexto)
    r = llm.chamar("redacao", s, u, max_tokens=6000, demo=demo)
    d = llm.extrair_json(r.texto)
    proposta.texto = str(d.get("texto") or "").strip()
    ia = [{"nivel": a.get("nivel") if a.get("nivel") in ("alto", "medio", "baixo") else "medio",
           "texto": str(a.get("texto") or "")[:600], "fundamento": str(a.get("fundamento") or "")[:200], "origem": "ia"}
          for a in d.get("alertas") or [] if a.get("texto")]
    proposta.alertas = ia + [a for a in (proposta.alertas or []) if a.get("origem") == "regra"]
    proposta.demonstracao = proposta.demonstracao or r.demonstracao
    return [r]


# ---------------------------------------------------------------- valor por extenso
_UN = ["", "um", "dois", "três", "quatro", "cinco", "seis", "sete", "oito", "nove", "dez", "onze", "doze", "treze",
       "quatorze", "quinze", "dezesseis", "dezessete", "dezoito", "dezenove"]
_DEZ = ["", "", "vinte", "trinta", "quarenta", "cinquenta", "sessenta", "setenta", "oitenta", "noventa"]
_CEM = ["", "cento", "duzentos", "trezentos", "quatrocentos", "quinhentos", "seiscentos", "setecentos", "oitocentos",
        "novecentos"]


def _ate_mil(n):
    if n == 100:
        return "cem"
    c, resto = divmod(n, 100)
    partes = [_CEM[c]] if c else []
    if resto < 20:
        if resto:
            partes.append(_UN[resto])
    else:
        d, u = divmod(resto, 10)
        partes.append(_DEZ[d] + (f" e {_UN[u]}" if u else ""))
    return " e ".join(p for p in partes if p)


def _inteiro(n):
    if n == 0:
        return "zero"
    grupos = []
    for nome_s, nome_p in (("", ""), ("mil", "mil"), ("milhão", "milhões"), ("bilhão", "bilhões")):
        n, g = divmod(n, 1000)
        grupos.append((g, nome_s, nome_p))
        if not n:
            break
    partes = []  # (texto, valor do grupo)
    for g, s, p in reversed(grupos):
        if not g:
            continue
        if s == "mil":
            partes.append(("mil" if g == 1 else f"{_ate_mil(g)} mil", g))
        elif s:
            partes.append((f"{_ate_mil(g)} {s if g == 1 else p}", g))
        else:
            partes.append((_ate_mil(g), g))
    # "dois milhões e quinhentos mil", "mil e cem", mas "dois mil trezentos e quarenta e cinco"
    txt = partes[0][0]
    for t, g in partes[1:]:
        txt += (" e " if g < 100 or g % 100 == 0 else " ") + t
    return txt


def extenso_reais(valor):
    valor = round(float(valor or 0), 2)
    reais, cent = int(valor), int(round((valor - int(valor)) * 100))
    txt = []
    if reais:
        prep = " de" if reais >= 1_000_000 and reais % 1_000_000 == 0 else ""
        txt.append(f"{_inteiro(reais)}{prep} {'real' if reais == 1 else 'reais'}")
    if cent:
        txt.append(f"{_inteiro(cent)} {'centavo' if cent == 1 else 'centavos'}")
    return " e ".join(txt) if txt else "zero real"


def moeda(v):
    if v is None:
        return "—"
    s = f"{float(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {s}"


# ---------------------------------------------------------------- Word
def docx(proposta, empresa):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2)
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)
    tot = totais(proposta)

    def texto_rico(par, linha):
        for i, parte in enumerate(re.split(r"\*\*", linha)):
            par.add_run(parte).bold = bool(i % 2)

    def tabela_itens():
        cols = ["Item", "Descrição", "Unid.", "Qtd.", "Preço unitário", "Preço total"]
        t = doc.add_table(rows=1, cols=len(cols))
        t.style = "Table Grid"
        for c, nome in zip(t.rows[0].cells, cols):
            c.text = ""
            c.paragraphs[0].add_run(nome).bold = True
        for n, it in enumerate(proposta.itens or [], start=1):
            q = it.get("quantidade")
            vals = [str(it.get("numero") or n), it.get("descricao") or "", it.get("unidade") or "",
                    (f"{q:g}" if isinstance(q, (int, float)) else str(q or "")).replace(".", ","),
                    moeda(it.get("preco_unitario")), moeda(it.get("preco_total"))]
            cells = t.add_row().cells
            for i, (c, v) in enumerate(zip(cells, vals)):
                c.text = v
                if i >= 3:
                    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        cells = t.add_row().cells
        cells[0].merge(cells[4]).text = ""
        cells[0].paragraphs[0].add_run("VALOR GLOBAL").bold = True
        cells[5].text = ""
        cells[5].paragraphs[0].add_run(moeda(tot["total"])).bold = True
        cells[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for row in t.rows:
            for i, c in enumerate(row.cells):
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9.5)
        doc.add_paragraph()

    valor_txt = f"{moeda(tot['total'])} ({extenso_reais(tot['total'])})"
    texto = proposta.texto or f"## Proposta comercial\n\n{{{{TABELA_ITENS}}}}\n\n**Valor global:** {{{{VALOR_TOTAL}}}}"
    tabela_feita = False
    for linha in texto.split("\n"):
        l = linha.rstrip()
        if "{{TABELA_ITENS}}" in l:
            tabela_itens()
            tabela_feita = True
            continue
        l = l.replace("{{VALOR_TOTAL}}", valor_txt)
        if not l.strip():
            continue
        if l.startswith("## ") or l.startswith("# "):
            h = doc.add_heading(l.lstrip("# ").strip(), level=2)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0x07, 0x1D, 0x2D)
        elif re.match(r"^\s*[-*] ", l):
            texto_rico(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", l))
        else:
            texto_rico(doc.add_paragraph(), l)
    if not tabela_feita:
        tabela_itens()
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def nome_arquivo(proposta):
    base = re.sub(r"[^\w\-]+", "_", (proposta.titulo or "proposta"))[:60].strip("_")
    return f"{base or 'proposta'}_{date.today().isoformat()}.docx"
